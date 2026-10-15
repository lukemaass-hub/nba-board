#!/usr/bin/env python3
"""Build backtest/results.xlsx (for Google Sheets) from data/results.json and data/dataset_stats.json."""
import datetime as dt, json, os, sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data")
R = json.load(open(os.path.join(DATA, "results.json")))
S = json.load(open(os.path.join(DATA, "dataset_stats.json")))
RUN = dt.date.today().isoformat()
SRC = (f"Source: backtest/backtest.py in github.com/lukemaass-hub/nba-board, on ESPN box scores "
       f"(2024-25 and 2025-26 regular seasons, {S['games']:,} games), run by Claude {RUN}.")
ORDER = ["Random", "Highest usual minutes", "Next man up (same position)", "Simple model", "Boosted ranker"]
HEAD = Font(name="Arial", bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name="Arial"); BOLD = Font(name="Arial", bold=True); TITLE = Font(name="Arial", bold=True, size=13)
PCT = "0.0%"; DEC = "0.000"; PTS = '+0.0%;-0.0%;"-"'
wb = Workbook()


def header(ws, row, cols, note=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row, i, c); cell.font = HEAD; cell.fill = HFILL; cell.alignment = Alignment(wrap_text=True, vertical="center")
    if note:
        ws.cell(row, 1).comment = Comment(note, "Claude")


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def style_body(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font != HEAD and not c.font.bold:
                c.font = BODY


# ---------------- Summary
ws = wb.active; ws.title = "Summary"
sp = R["splits"]["Train 2024-25, test 2025-26"]
ws["A1"] = (f"Both models beat every simple rule: the #1 pick gained minutes {100*sp['Simple model']['top1']:.0f}% of the time, "
            f"vs {100*sp['Next man up (same position)']['top1']:.0f}% for next man up"); ws["A1"].font = TITLE
ws["A2"] = ("Ranking of healthy teammates by chance of playing 2+ minutes above their 10-game baseline, "
            "in team-games where a rotation player (15+ min baseline) is out.")
r = 4
for label, block in list(R["splits"].items()) + [("Walk-forward through 2025-26 (monthly re-train)", R["walk_forward"])]:
    ws.cell(r, 1, label).font = BOLD; r += 1
    header(ws, r, ["Method", "Units", "#1 pick is a gainer", "Top 3 precision", "Rank correlation",
                   "#1 pick vs best simple rule", "Top 3 vs best simple rule"], SRC)
    first = r + 1
    for k in ORDER:
        r += 1
        v = block[k]
        ws.cell(r, 1, k); ws.cell(r, 2, v["units"]); ws.cell(r, 3, v["top1"]).number_format = PCT
        ws.cell(r, 4, v["p3"]).number_format = PCT; ws.cell(r, 5, v["rho"]).number_format = DEC
    # best simple rule = the better of "Highest usual minutes" and "Next man up" (rows first+1, first+2)
    for rr in (first + 3, first + 4):
        ws.cell(rr, 6, f"=C{rr}-MAX(C{first + 1}:C{first + 2})").number_format = PTS
        ws.cell(rr, 7, f"=D{rr}-MAX(D{first + 1}:D{first + 2})").number_format = PTS
    r += 1
    if "_vs_best_baseline" in block:
        for m, d in block["_vs_best_baseline"].items():
            ws.cell(r, 1, f"{m} vs {d['baseline']}, 95% range of the #1-pick gap: {100*d['top1'][0]:+.1f} to {100*d['top1'][1]:+.1f} points "
                          f"(if both ends are above 0, the model is better)"); r += 1
    r += 1
ws.cell(r, 1, f"Gainer = played 2+ minutes more than his average of the previous 10 games. "
              f"{100*S['gainer_rate']:.1f}% of healthy teammates were gainers, so a random pick scores about that.")
widths(ws, [34, 9, 15, 13, 13, 17, 15]); ws.freeze_panes = "A4"
for row in ws.iter_rows(min_row=3):
    for c in row:
        if c.row > 2:
            c.alignment = Alignment(wrap_text=c.column > 1 or c.font == HEAD, vertical="center")

# ---------------- Walk-forward by month
ws = wb.create_sheet("By month")
ws["A1"] = "Walk-forward through 2025-26: #1 pick is a gainer, by month"; ws["A1"].font = TITLE
months = sorted({k.split("|")[0] for k in R["walk_forward_by_month"]})
header(ws, 3, ["Month", "Units"] + ORDER, SRC)
for i, mth in enumerate(months, 4):
    ws.cell(i, 1, f"'{mth}"); ws.cell(i, 2, R["walk_forward_by_month"][f"{mth}|Simple model"]["units"])
    for j, k in enumerate(ORDER, 3):
        ws.cell(i, j, R["walk_forward_by_month"][f"{mth}|{k}"]["top1"]).number_format = PCT
last = 3 + len(months)
ws.cell(last + 1, 1, "Weighted by units").font = BOLD
ws.cell(last + 1, 2, f"=SUM(B4:B{last})")
for j in range(3, 3 + len(ORDER)):
    col = get_column_letter(j)
    c = ws.cell(last + 1, j, f"=SUMPRODUCT($B4:$B{last},{col}4:{col}{last})/$B{last + 1}"); c.number_format = PCT; c.font = BOLD
widths(ws, [12, 9] + [16] * len(ORDER)); ws.freeze_panes = "A4"

# ---------------- What the models look at
ws = wb.create_sheet("What matters")
ws["A1"] = "What the models look at"; ws["A1"].font = TITLE
header(ws, 3, ["Simple model input", "Weight (+ = more likely to gain)"], SRC)
for i, (n, w) in enumerate(R["simple_weights"], 4):
    ws.cell(i, 1, n); ws.cell(i, 2, w).number_format = "+0.00;-0.00"
r0 = 5 + len(R["simple_weights"])
header(ws, r0, ["Boosted ranker input", "Share of what it learned"])
for i, (n, w) in enumerate(R["gbm_importance"], r0 + 1):
    ws.cell(i, 1, n); ws.cell(i, 2, w / 100).number_format = PCT
ws.cell(r0 + 2 + len(R["gbm_importance"]), 1, "Weights are on a common scale (each input standardized). Both models were trained on both seasons for this tab.")
widths(ws, [58, 30])

# ---------------- Where it fails
ws = wb.create_sheet("Where it fails")
ws["A1"] = f"Where the model's #1 pick goes wrong ({R['best_model']}, tested on 2025-26)"; ws["A1"].font = TITLE
r = 3
for k, d in R["fail_by"].items():
    if k == "Units":
        for kk, vv in d.items():
            ws.cell(r, 1, kk); ws.cell(r, 2, vv).number_format = PCT; r += 1
        r += 1; continue
    ws.cell(r, 1, k).font = BOLD; r += 1
    header(ws, r, ["Group", "Model #1 pick right", "Next man up #1 pick right", "Model minus next man up"], SRC if r == 5 else None)
    for kk, vv in d.items():
        r += 1
        ws.cell(r, 1, f"'{kk}"); ws.cell(r, 2, vv["top1"]).number_format = PCT; ws.cell(r, 3, vv["nextman_top1"]).number_format = PCT
        ws.cell(r, 4, f"=B{r}-C{r}").number_format = PTS
    r += 2
widths(ws, [44, 20, 24, 22])

# ---------------- Examples
ws = wb.create_sheet("Examples")
ws["A1"] = "Examples from 2025-26 (games with 30+ minutes missing)"; ws["A1"].font = TITLE
r = 3
for title, key in (("Hits: the #1 pick gained minutes", "examples_hit"), ("Misses: the #1 pick did not, though someone else did", "examples_miss")):
    ws.cell(r, 1, title).font = BOLD; r += 1
    header(ws, r, ["Date", "Team", "Opponent", "Out (baseline minutes)", "Model's top 3 (actual change)", "Biggest real gainer"], SRC)
    for e in R[key]:
        r += 1
        for j, c in enumerate(["date", "team", "opponent", "out_players", "model_top3", "real_biggest"], 1):
            ws.cell(r, j, e[c]).alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
widths(ws, [12, 8, 10, 40, 48, 28])

# ---------------- Data Sources
ws = wb.create_sheet("Data Sources")
cols = ["Written to", "Source", "Object", "Obtained via", "Query or request", "Parameters", "Run at", "Requested by",
        "Row count", "Changes after retrieval", "Notes"]
header(ws, 1, cols)
ws.append(["Summary, By month, What matters, Where it fails, Examples", "ESPN public JSON (scoreboard + game summary box scores)",
           "site.api.espn.com/apis/site/v2/sports/basketball/nba/summary?event=ID", "Downloaded and scored by Claude (backtest/ scripts)",
           "Backtest which healthy teammates gain 2+ minutes when a rotation teammate is out; train 2024-25 test 2025-26, swap, walk-forward",
           "Regular season only; players on more than one team dropped; out = inactive or DNP injury/illness/rest; baseline = previous 10 games played",
           RUN, "Luke", f"{S['units']:,} test units, {S['candidates']:,} teammate rows",
           "Aggregated by backtest/backtest.py; ESPN does not list most inactive players, so 'missing while on the team' counts as inactive",
           "Full report: backtest/REPORT.md"])
widths(ws, [34, 30, 40, 26, 50, 50, 12, 12, 22, 50, 24])
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

for w in wb.worksheets:
    style_body(w)
out = os.path.join(HERE, "results.xlsx") if len(sys.argv) <= 1 else os.path.join(DATA, "results.xlsx")
wb.save(out); print("wrote", out, os.path.getsize(out), "bytes")
